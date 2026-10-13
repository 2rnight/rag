"""多格式文档解析（混合解析器）。

设计思路（面试可讲）：
1. **统一解析接口**：`BaseParser.parse() -> list[ParsedDocument]`，上层（切分/入库）
   完全不感知格式差异，新增格式只需注册一个 Parser。
2. **结构化优先**：Markdown / 纯文本目录类文档先做“结构探测”，识别到标题层级就
   按章节成块（`heading_path` 一路带进 metadata 与引用来源），避免整篇文本被切碎后丢失语义。
3. **PDF 双策略竞争**：同时尝试 pypdf 的普通模式与 layout 模式，按“文本质量分”
   （有效字符占比、乱码率、碎片行占比）择优；再做页眉页脚去重、断行粘连修复，
   识别不到文字层时给出“疑似扫描件，需要 OCR”的明确报错，而不是静默产出垃圾数据。
4. **可观测**：每个文档记录使用的解析器、质量分、页数，随入库报告一起返回。
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from abc import ABC, abstractmethod
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .exceptions import DocumentParseError, EmptyDocumentError, UnsupportedFileTypeError
from .logging_utils import get_logger, stopwatch

logger = get_logger(__name__)

# ---------------------------------------------------------------- 数据结构


@dataclass(slots=True)
class ParsedDocument:
    """解析后的“逻辑块”：一段语义连续的文本 + 结构元数据。"""

    text: str
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def title(self) -> str:
        return str(self.metadata.get("heading_path") or self.metadata.get("source_file") or "")


@dataclass(slots=True)
class ParseReport:
    """单个文件的解析报告。"""

    source_file: str
    doc_type: str
    parser: str
    blocks: int
    chars: int
    quality: float
    elapsed_ms: float
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- 文本工具

_CJK = r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
_CJK_RE = re.compile(f"[{_CJK}]")
_REPLACEMENT_CHARS = ("\ufffd", "\x00", "\u000b")
_PAGE_NUMBER_RE = re.compile(r"^\s*[-–—\[(]?\s*\d{1,4}\s*[-–—\])]?\s*$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_LIST_OR_STRUCT_LINE = re.compile(
    r"^\s*(?:[-*+•·]|\d+[.、)]|[（(]\d+[)）]|#{1,6}\s|\||>|\w+\s*[:：]\s|```)"
)
_SENTENCE_END = "。！？；!?;:："


def text_quality(text: str) -> float:
    """粗粒度文本质量分：有效字符（中日韩/字母数字）占比，惩罚乱码与碎片行。"""
    if not text or not text.strip():
        return 0.0
    total = len(text)
    valid = len(_CJK_RE.findall(text)) + sum(ch.isalnum() for ch in text)
    bad = sum(text.count(ch) for ch in _REPLACEMENT_CHARS)
    lines = [ln for ln in text.splitlines() if ln.strip()]
    fragment_ratio = (
        sum(1 for ln in lines if len(ln.strip()) <= 2) / len(lines) if lines else 1.0
    )
    score = valid / total - 3.0 * bad / total - 0.3 * fragment_ratio
    return round(max(0.0, score), 4)


def decode_bytes(data: bytes) -> tuple[str, str]:
    """多编码探测解码（中文文档常见 UTF-8 / GB18030）。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8(replace)"


def normalize_text(text: str) -> str:
    """统一换行、去掉零宽字符、压缩多余空行。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u200b", "").replace("\ufeff", "").replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def join_wrapped_lines(text: str) -> str:
    """把被 PDF 硬换行拆断的句子接回去（中英文分别处理）。"""
    lines = text.split("\n")
    out: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not out or not stripped:
            out.append(line)
            continue
        prev = out[-1]
        prev_stripped = prev.strip()
        if (
            prev_stripped
            and not _LIST_OR_STRUCT_LINE.match(prev_stripped)
            and not _LIST_OR_STRUCT_LINE.match(stripped)
            and prev_stripped[-1] not in _SENTENCE_END
            and not prev_stripped.endswith(".")
        ):
            # 英文连字符断行：mer-ging -> merging
            if prev_stripped.endswith("-") and stripped[:1].isascii():
                out[-1] = prev_stripped[:-1] + stripped
                continue
            need_space = prev_stripped[-1].isascii() and stripped[0].isascii()
            out[-1] = prev_stripped + (" " if need_space else "") + stripped
            continue
        out.append(line)
    return "\n".join(out)


def strip_repeated_headers_footers(pages: list[str], threshold: float = 0.6) -> list[str]:
    """去掉在多页重复出现的页眉/页脚（例如公司名、页码、"机密"）。"""
    if len(pages) < 3:
        return pages
    counter: dict[str, int] = {}
    for page in pages:
        seen = {
            ln.strip()
            for ln in page.split("\n")[:2] + page.split("\n")[-2:]
            if 0 < len(ln.strip()) <= 60
        }
        for ln in seen:
            counter[ln] = counter.get(ln, 0) + 1
    repeated = {ln for ln, c in counter.items() if c / len(pages) >= threshold}
    if not repeated:
        return pages
    cleaned = []
    for page in pages:
        kept = [ln for ln in page.split("\n") if ln.strip() not in repeated]
        cleaned.append("\n".join(kept))
    return cleaned


def _clean_page_text(text: str) -> str:
    lines = [ln for ln in text.split("\n") if not _PAGE_NUMBER_RE.match(ln)]
    return join_wrapped_lines(normalize_text("\n".join(lines)))


# ---------------------------------------------------------------- 解析器基类


class BaseParser(ABC):
    """解析器统一接口。"""

    name: str = "base"
    suffixes: tuple[str, ...] = ()
    doc_type: str = "text"

    @abstractmethod
    def parse(self, path: Path) -> list[ParsedDocument]:
        """把文件解析成若干语义块（块内部再交给切分器做重叠切片）。"""


# ---------------------------------------------------------------- 结构化文本


def split_by_headings(text: str, min_section_chars: int = 40) -> list[tuple[str, str, list[str]]]:
    """按 Markdown 标题层级切分。

    返回 [(heading_path, section_text, heading_paths), ...]；
    合并过短小节时，被合并小节的标题会保留在 heading_paths 里（引用溯源不丢信息）。
    支持 `#` 标题与代码围栏（围栏内的 `#` 不当作标题）。
    """
    sections: list[tuple[str, str, list[str]]] = []
    stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    in_fence = False

    def flush() -> None:
        content = normalize_text("\n".join(buffer))
        if content:
            path = " > ".join(title for _, title in stack) or "正文"
            sections.append((path, content, [path]))
        buffer.clear()

    for line in text.split("\n"):
        if line.strip().startswith("```"):
            in_fence = not in_fence
            buffer.append(line)
            continue
        match = None if in_fence else _HEADING_RE.match(line)
        if match:
            flush()
            level, title = len(match.group(1)), match.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            continue
        buffer.append(line)
    flush()

    # 合并过短的相邻小节，避免制造大量碎片块（被合并小节的小标题记进 heading_paths）
    merged: list[tuple[str, str, list[str]]] = []
    for path, content, paths in sections:
        if (
            merged
            and len(content) < min_section_chars
            and len(merged[-1][1]) + len(content) <= 1500
        ):
            prev_path, prev_content, prev_paths = merged[-1]
            extra = [item for item in paths if item != prev_path]
            merged[-1] = (prev_path, f"{prev_content}\n{content}", prev_paths + extra)
        else:
            merged.append((path, content, paths))
    return merged


class MarkdownParser(BaseParser):
    """Markdown 解析器：保留标题层级，长小节交给下游切分器再切。"""

    name = "markdown"
    suffixes = (".md", ".markdown", ".mdx")
    doc_type = "markdown"

    def __init__(self, max_block_chars: int = 2000) -> None:
        self.max_block_chars = max_block_chars

    def parse(self, path: Path) -> list[ParsedDocument]:
        raw, _ = decode_bytes(path.read_bytes())
        return self.parse_text(raw, heading_path_prefix="")

    def parse_text(self, raw: str, heading_path_prefix: str = "") -> list[ParsedDocument]:
        text = normalize_text(raw)
        # 去掉 YAML front matter，但把标题提取出来
        front_matter_title = ""
        if text.startswith("---"):
            end = text.find("\n---", 3)
            if end != -1:
                front = text[3:end]
                text = text[end + 4 :]
                for line in front.split("\n"):
                    if line.strip().lower().startswith("title:"):
                        front_matter_title = line.split(":", 1)[1].strip().strip("\"'")

        docs: list[ParsedDocument] = []
        for path_str, content, path_list in split_by_headings(text):
            heading_path = f"{heading_path_prefix}{path_str}" if heading_path_prefix else path_str
            for block in self._split_long(content):
                docs.append(
                    ParsedDocument(
                        text=block,
                        metadata={
                            "heading_path": heading_path,
                            "heading_paths": list(path_list),
                            "doc_type": self.doc_type,
                            "parser": self.name,
                            "title": front_matter_title or heading_path.split(" > ")[-1],
                        },
                    )
                )
        return docs

    def _split_long(self, content: str) -> list[str]:
        if len(content) <= self.max_block_chars:
            return [content]
        blocks: list[str] = []
        current: list[str] = []
        size = 0
        for para in content.split("\n\n"):
            if size + len(para) > self.max_block_chars and current:
                blocks.append("\n\n".join(current))
                current, size = [], 0
            current.append(para)
            size += len(para) + 2
        if current:
            blocks.append("\n\n".join(current))
        return blocks


class TextParser(BaseParser):
    """纯文本解析器：探测到标题结构就按结构化解析，否则按空行聚合段落。"""

    name = "text"
    suffixes = (".txt", ".text", ".log")
    doc_type = "text"

    def __init__(self, max_block_chars: int = 2000, structure_probe: bool = True) -> None:
        self.max_block_chars = max_block_chars
        self.structure_probe = structure_probe
        self._markdown = MarkdownParser(max_block_chars=max_block_chars)

    def parse(self, path: Path) -> list[ParsedDocument]:
        raw, encoding = decode_bytes(path.read_bytes())
        text = normalize_text(raw)
        heading_lines = [ln for ln in text.split("\n") if _HEADING_RE.match(ln)]
        use_structure = self.structure_probe and len(heading_lines) >= 2

        if use_structure:
            docs = self._markdown.parse_text(text)
            for doc in docs:
                doc.metadata.update({"doc_type": self.doc_type, "parser": f"{self.name}+heading"})
            return docs

        docs: list[ParsedDocument] = []
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        current: list[str] = []
        size = 0
        for para in paragraphs:
            if size + len(para) > self.max_block_chars and current:
                docs.append(self._make_doc("\n\n".join(current), encoding))
                current, size = [], 0
            current.append(para)
            size += len(para) + 2
        if current:
            docs.append(self._make_doc("\n\n".join(current), encoding))
        return docs

    def _make_doc(self, text: str, encoding: str) -> ParsedDocument:
        return ParsedDocument(
            text=text,
            metadata={"doc_type": self.doc_type, "parser": self.name, "encoding": encoding},
        )


class CsvParser(BaseParser):
    """CSV/TSV 解析器：每行转成 `列名: 值` 文本，按行分组避免过碎。"""

    name = "csv"
    suffixes = (".csv", ".tsv")
    doc_type = "table"

    def __init__(self, rows_per_block: int = 20) -> None:
        self.rows_per_block = rows_per_block

    def parse(self, path: Path) -> list[ParsedDocument]:
        raw, encoding = decode_bytes(path.read_bytes())
        delimiter = "\t" if path.suffix.lower() == ".tsv" else self._sniff(raw)
        reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
        rows = [row for row in reader if any(cell.strip() for cell in row)]
        if not rows:
            raise EmptyDocumentError(path.name, "CSV 中没有数据行")

        header = [h.strip() or f"列{i + 1}" for i, h in enumerate(rows[0])]
        docs: list[ParsedDocument] = []
        data_rows = rows[1:] if len(rows) > 1 else []
        for start in range(0, len(data_rows), self.rows_per_block):
            block_rows = data_rows[start : start + self.rows_per_block]
            lines = []
            for row in block_rows:
                pairs = [
                    f"{header[i] if i < len(header) else f'列{i + 1}'}: {cell.strip()}"
                    for i, cell in enumerate(row)
                    if cell.strip()
                ]
                lines.append(" | ".join(pairs))
            docs.append(
                ParsedDocument(
                    text=f"表头: {' | '.join(header)}\n" + "\n".join(lines),
                    metadata={
                        "doc_type": self.doc_type,
                        "parser": self.name,
                        "heading_path": f"表格 第{start + 1}-{start + len(block_rows)}行",
                        "row_start": start + 1,
                        "row_end": start + len(block_rows),
                        "encoding": encoding,
                    },
                )
            )
        return docs

    @staticmethod
    def _sniff(raw: str) -> str:
        sample = "\n".join(raw.split("\n")[:20])
        try:
            return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except Exception:  # noqa: BLE001 - 探测失败就退回逗号
            return ","


class JsonParser(BaseParser):
    """JSON 解析器：列表型数据按条目拆分，方便检索到具体记录。"""

    name = "json"
    suffixes = (".json",)
    doc_type = "json"

    def __init__(self, items_per_block: int = 10) -> None:
        self.items_per_block = items_per_block

    def parse(self, path: Path) -> list[ParsedDocument]:
        raw, _ = decode_bytes(path.read_bytes())
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise DocumentParseError(f"JSON 解析失败: {exc}") from exc

        if isinstance(payload, list):
            docs: list[ParsedDocument] = []
            for start in range(0, len(payload), self.items_per_block):
                items = payload[start : start + self.items_per_block]
                text = "\n".join(json.dumps(item, ensure_ascii=False) for item in items)
                docs.append(
                    ParsedDocument(
                        text=text,
                        metadata={
                            "doc_type": self.doc_type,
                            "parser": self.name,
                            "heading_path": f"条目 {start + 1}-{start + len(items)}",
                        },
                    )
                )
            return docs or [ParsedDocument(text=raw, metadata={"doc_type": self.doc_type, "parser": self.name})]

        pretty = json.dumps(payload, ensure_ascii=False, indent=2)
        return [
            ParsedDocument(
                text=pretty,
                metadata={"doc_type": self.doc_type, "parser": self.name, "heading_path": "JSON 根对象"},
            )
        ]


class DocxParser(BaseParser):
    """Word(.docx) 解析器：段落 + 标题样式 + 表格（转 Markdown 表格）。"""

    name = "docx"
    suffixes = (".docx",)
    doc_type = "docx"

    def parse(self, path: Path) -> list[ParsedDocument]:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError("未安装 python-docx，无法解析 .docx") from exc

        document = docx.Document(str(path))
        sections: list[tuple[str, list[str]]] = []
        stack: list[tuple[int, str]] = []
        buffer: list[str] = []

        def flush() -> None:
            content = normalize_text("\n".join(buffer))
            if content:
                sections.append((" > ".join(t for _, t in stack) or "正文", content))
            buffer.clear()

        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style else ""
            level = None
            if style.startswith("heading"):
                digits = re.findall(r"\d+", style)
                level = int(digits[0]) if digits else 1
            elif style in {"title", "subtitle"}:
                level = 1
            if level:
                flush()
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
                continue
            buffer.append(text)
        flush()

        for table in document.tables:
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if rows:
                sections.append(("表格", "\n".join(rows)))

        if not sections:
            raise EmptyDocumentError(path.name, "文档没有可提取的段落文本")

        return [
            ParsedDocument(
                text=content,
                metadata={"doc_type": self.doc_type, "parser": self.name, "heading_path": heading},
            )
            for heading, content in sections
        ]


class HtmlParser(BaseParser):
    """HTML 解析器：优先正文容器，转成 Markdown 风格的标题/列表/表格。"""

    name = "html"
    suffixes = (".html", ".htm", ".xhtml")
    doc_type = "html"

    def parse(self, path: Path) -> list[ParsedDocument]:
        raw, encoding = decode_bytes(path.read_bytes())
        return self.parse_html(raw, encoding=encoding, source=path.name)

    def parse_html(self, raw: str, encoding: str = "utf-8", source: str = "<html>") -> list[ParsedDocument]:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError("未安装 beautifulsoup4，无法解析 HTML") from exc

        soup = BeautifulSoup(raw, "lxml" if _has_lxml() else "html.parser")
        for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form"]):
            tag.decompose()

        container = soup.find("main") or soup.find("article") or soup.body or soup

        md_lines: list[str] = []
        for element in container.descendants:
            name = getattr(element, "name", None)
            if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                level = int(name[1])
                text = element.get_text(" ", strip=True)
                if text:
                    md_lines.append(f"\n{'#' * level} {text}")
            elif name == "p":
                text = element.get_text(" ", strip=True)
                if text:
                    md_lines.append(text)
            elif name == "li":
                text = element.get_text(" ", strip=True)
                if text:
                    md_lines.append(f"- {text}")
            elif name == "table":
                table_lines = []
                for row in element.find_all("tr"):
                    cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["td", "th"])]
                    if any(cells):
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    md_lines.append("\n".join(table_lines))

        text = normalize_text("\n".join(md_lines))
        if not text:
            raise EmptyDocumentError(source, "页面没有可提取的正文（可能是纯 JS 渲染页面）")

        markdown = MarkdownParser()
        docs = markdown.parse_text(text)
        for doc in docs:
            doc.metadata.update({"doc_type": self.doc_type, "parser": self.name, "encoding": encoding})
        return docs


def _has_lxml() -> bool:
    try:
        import lxml  # noqa: F401

        return True
    except ImportError:  # pragma: no cover
        return False


class PdfParser(BaseParser):
    """PDF 解析器：双策略竞争（普通 / layout）+ 乱码清理 + 扫描件识别。"""

    name = "pdf"
    suffixes = (".pdf",)
    doc_type = "pdf"
    strategies = ("pypdf", "pypdf_layout")

    def __init__(self, min_quality: float = 0.25) -> None:
        self.min_quality = min_quality

    def parse(self, path: Path) -> list[ParsedDocument]:
        candidates: list[tuple[str, list[str], float]] = []
        errors: list[str] = []

        for strategy in self.strategies:
            try:
                pages = self._extract(path, layout=strategy.endswith("layout"))
            except Exception as exc:  # noqa: BLE001 - 单个策略失败不影响另一策略
                errors.append(f"{strategy}: {exc}")
                continue
            if not pages:
                errors.append(f"{strategy}: 未提取到文本")
                continue
            joined = "\n".join(pages)
            candidates.append((strategy, pages, text_quality(joined)))

        if not candidates:
            raise EmptyDocumentError(
                path.name,
                "PDF 无文字层，疑似扫描件，需要 OCR（可先转成文本/Markdown 再入库）。"
                + (f" 底层错误: {'; '.join(errors)}" if errors else ""),
            )

        strategy, pages, quality = max(candidates, key=lambda item: item[2])
        if quality < self.min_quality:
            raise EmptyDocumentError(
                path.name,
                f"PDF 文本质量过低（quality={quality}），可能是扫描件或字体未内嵌，建议走 OCR。",
            )

        pages = strip_repeated_headers_footers([_clean_page_text(p) for p in pages])

        docs: list[ParsedDocument] = []
        for index, page_text in enumerate(pages):
            if len(page_text) < 10:
                continue
            docs.append(
                ParsedDocument(
                    text=page_text,
                    metadata={
                        "doc_type": self.doc_type,
                        "parser": f"{self.name}:{strategy}",
                        "page": index,  # 0-based，对外输出时 +1
                        "heading_path": f"第 {index + 1} 页",
                        "page_quality": text_quality(page_text),
                    },
                )
            )
        if not docs:
            raise EmptyDocumentError(path.name, "PDF 每页文本都过短")
        logger.info(
            "PDF 解析完成",
            extra={"extra_fields": {"file": path.name, "strategy": strategy, "quality": quality, "pages": len(docs)}},
        )
        return docs

    @staticmethod
    def _extract(path: Path, layout: bool = False) -> list[str]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")  # 空密码（仅限制打印/复制的 PDF）
            except Exception as exc:  # noqa: BLE001
                raise DocumentParseError(f"PDF 已加密，无法解析: {exc}") from exc

        pages: list[str] = []
        for page in reader.pages:
            if layout:
                try:
                    pages.append(page.extract_text(extraction_mode="layout") or "")
                except TypeError:  # 旧版 pypdf 不支持 layout 模式
                    pages.append(page.extract_text() or "")
            else:
                pages.append(page.extract_text() or "")
        return pages


# ---------------------------------------------------------------- 统一加载器


class HybridDocumentLoader:
    """混合文档加载器：按扩展名路由到具体解析器，并补齐文件级元数据。"""

    def __init__(self, parsers: Iterable[BaseParser] | None = None) -> None:
        parser_list = list(parsers) if parsers else [
            PdfParser(),
            MarkdownParser(),
            TextParser(),
            DocxParser(),
            CsvParser(),
            JsonParser(),
            HtmlParser(),
        ]
        self._parsers: dict[str, BaseParser] = {}
        for parser in parser_list:
            for suffix in parser.suffixes:
                self._parsers[suffix] = parser
        self.last_report: list[ParseReport] = []

    @property
    def supported_suffixes(self) -> list[str]:
        return sorted(self._parsers)

    def supports(self, path: str | Path) -> bool:
        return Path(path).suffix.lower() in self._parsers

    # -------------------------------------------------- 对外接口
    def parse(self, path: str | Path, *, root: Path | None = None) -> list[ParsedDocument]:
        file_path = Path(path)
        if not file_path.exists():
            raise DocumentParseError(f"文件不存在: {file_path}")
        suffix = file_path.suffix.lower()
        parser = self._parsers.get(suffix)
        if parser is None:
            raise UnsupportedFileTypeError(suffix, self.supported_suffixes)

        raw_bytes = file_path.read_bytes()
        file_hash = hashlib.sha256(raw_bytes).hexdigest()
        try:
            relative = str(file_path.resolve().relative_to(root.resolve())) if root else file_path.name
        except ValueError:
            relative = file_path.name

        with stopwatch() as sw:
            docs = parser.parse(file_path)

        warnings: list[str] = []
        kept: list[ParsedDocument] = []
        for doc in docs:
            text = normalize_text(doc.text)
            if len(text) < 5:
                continue
            doc.text = text
            doc.metadata.setdefault("doc_type", parser.doc_type)
            doc.metadata.setdefault("parser", parser.name)
            doc.metadata.update(
                {
                    "source_file": file_path.name,
                    "source_path": relative,
                    "file_hash": file_hash,
                    "size_bytes": len(raw_bytes),
                    "ingested_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                }
            )
            kept.append(doc)

        quality = text_quality("\n".join(d.text for d in kept))
        if quality < 0.3:
            warnings.append(f"文本质量偏低（quality={quality}），检索效果可能受影响")
        if not kept:
            raise EmptyDocumentError(file_path.name)

        self.last_report = [
            ParseReport(
                source_file=file_path.name,
                doc_type=parser.doc_type,
                parser=parser.name,
                blocks=len(kept),
                chars=sum(len(d.text) for d in kept),
                quality=quality,
                elapsed_ms=round(sw["ms"], 2),
                warnings=warnings,
            )
        ]
        logger.info(
            "文档解析完成",
            extra={
                "extra_fields": {
                    "file": file_path.name,
                    "parser": parser.name,
                    "blocks": len(kept),
                    "chars": sum(len(d.text) for d in kept),
                    "quality": quality,
                    "ms": round(sw["ms"], 2),
                }
            },
        )
        return kept

    def parse_directory(self, directory: str | Path, *, recursive: bool = True) -> list[ParsedDocument]:
        root = Path(directory)
        if not root.exists():
            raise DocumentParseError(f"目录不存在: {root}")
        pattern = "**/*" if recursive else "*"
        all_docs: list[ParsedDocument] = []
        for file_path in sorted(root.glob(pattern)):
            if not file_path.is_file() or not self.supports(file_path):
                continue
            try:
                all_docs.extend(self.parse(file_path, root=root))
            except EmptyDocumentError as exc:
                logger.warning("跳过文件 %s: %s", file_path.name, exc)
            except DocumentParseError as exc:
                logger.warning("解析失败 %s: %s", file_path.name, exc)
        return all_docs
